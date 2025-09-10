"""
FileConnector 测试套件
测试CSV和Excel文件连接器的功能、性能和错误处理
"""
import pytest
import asyncio
import tempfile
import os
import time
import csv
from pathlib import Path
from typing import Dict, Any, List
import polars as pl
from io import StringIO
import openpyxl

# 测试目标导入
from connectors.file_connector import FileConnector, CSVProcessor, ExcelProcessor
from connectors.base import DataSourceType, QueryOptions, ConnectionStatus
from connectors.registry import registry
from utils.encoding_detector import encoding_detector
from utils.file_utils import file_utils
from core.memory_guard import memory_guard
from core.config import config


class TestFileUtils:
    """文件工具测试"""
    
    def test_get_file_type_csv(self, tmp_path):
        """测试CSV文件类型检测"""
        # 创建测试CSV文件
        csv_file = tmp_path / "test.csv"
        csv_file.write_text("name,age,city\nJohn,25,NYC\nJane,30,LA")
        
        file_type = file_utils.get_file_type(str(csv_file))
        assert file_type == 'csv'
    
    def test_get_file_type_excel(self, tmp_path):
        """测试Excel文件类型检测"""
        # 创建测试Excel文件
        excel_file = tmp_path / "test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['name', 'age', 'city'])
        ws.append(['John', 25, 'NYC'])
        wb.save(str(excel_file))
        
        file_type = file_utils.get_file_type(str(excel_file))
        assert file_type == 'excel'
    
    def test_detect_csv_delimiter(self, tmp_path):
        """测试CSV分隔符检测"""
        # 测试逗号分隔
        csv_file = tmp_path / "comma.csv"
        csv_file.write_text("name,age,city\nJohn,25,NYC")
        delimiter = file_utils.detect_csv_delimiter(str(csv_file))
        assert delimiter == ','
        
        # 测试制表符分隔
        tsv_file = tmp_path / "tab.tsv"
        tsv_file.write_text("name\tage\tcity\nJohn\t25\tNYC")
        delimiter = file_utils.detect_csv_delimiter(str(tsv_file))
        assert delimiter == '\t'
        
        # 测试分号分隔
        semi_file = tmp_path / "semi.csv"
        semi_file.write_text("name;age;city\nJohn;25;NYC")
        delimiter = file_utils.detect_csv_delimiter(str(semi_file))
        assert delimiter == ';'
    
    def test_estimate_file_info(self, tmp_path):
        """测试文件信息估算"""
        # 创建测试文件
        csv_file = tmp_path / "test.csv"
        data = []
        for i in range(1000):
            data.append(f"Name{i},Age{i},City{i}")
        csv_file.write_text("\n".join(data))
        
        file_info = file_utils.estimate_file_info(str(csv_file))
        
        assert file_info['file_type'] == 'csv'
        assert file_info['size_bytes'] > 0
        assert file_info['estimated_rows'] > 0
        assert file_info['delimiter'] == ','
    
    def test_validate_file_access(self, tmp_path):
        """测试文件访问验证"""
        # 测试存在的文件
        csv_file = tmp_path / "test.csv"
        csv_file.write_text("name,age\nJohn,25")
        
        result = file_utils.validate_file_access(str(csv_file))
        assert result['exists'] is True
        assert result['readable'] is True
        assert result['size_ok'] is True
        assert result['type_supported'] is True
        
        # 测试不存在的文件
        result = file_utils.validate_file_access("/non/existent/file.csv")
        assert result['exists'] is False
        assert len(result['errors']) > 0


class TestEncodingDetector:
    """编码检测器测试"""
    
    def test_detect_utf8_encoding(self, tmp_path):
        """测试UTF-8编码检测"""
        file_path = tmp_path / "utf8.csv"
        content = "姓名,年龄,城市\n张三,25,北京\n李四,30,上海"
        file_path.write_text(content, encoding='utf-8')
        
        result = encoding_detector.detect_encoding(str(file_path))
        assert result['encoding'] in ['utf-8', 'utf-8-sig']
        assert result['confidence'] > 0.7
    
    def test_detect_gbk_encoding(self, tmp_path):
        """测试GBK编码检测"""
        file_path = tmp_path / "gbk.csv"
        content = "姓名,年龄,城市\n张三,25,北京\n李四,30,上海"
        
        with open(str(file_path), 'w', encoding='gbk') as f:
            f.write(content)
        
        result = encoding_detector.detect_encoding(str(file_path))
        # GBK文件应该被检测为GBK或相关编码
        assert result['encoding'] in ['gbk', 'gb2312', 'gb18030']
    
    def test_detect_bom(self, tmp_path):
        """测试BOM检测"""
        file_path = tmp_path / "utf8_bom.csv"
        content = "name,age\nJohn,25"
        
        with open(str(file_path), 'w', encoding='utf-8-sig') as f:
            f.write(content)
        
        result = encoding_detector.detect_encoding(str(file_path))
        assert result['encoding'] == 'utf-8-sig'
        assert result['method'] == 'bom'
    
    def test_validate_encoding(self, tmp_path):
        """测试编码验证"""
        file_path = tmp_path / "test.csv"
        file_path.write_text("name,age\nJohn,25", encoding='utf-8')
        
        # 正确编码应该验证通过
        assert encoding_detector.validate_encoding(str(file_path), 'utf-8') is True
        
        # 错误编码应该验证失败
        assert encoding_detector.validate_encoding(str(file_path), 'gbk') is False


class TestCSVProcessor:
    """CSV处理器测试"""
    
    @pytest.fixture
    def sample_csv_file(self, tmp_path):
        """创建样本CSV文件"""
        csv_file = tmp_path / "sample.csv"
        data = []
        data.append("name,age,city,salary")
        for i in range(1000):
            data.append(f"Name{i},{20 + i % 50},City{i % 10},{30000 + i * 100}")
        csv_file.write_text("\n".join(data))
        return str(csv_file)
    
    @pytest.mark.asyncio
    async def test_csv_read_full(self, sample_csv_file):
        """测试完整CSV读取"""
        processor = CSVProcessor(sample_csv_file, encoding='utf-8', delimiter=',')
        df = await processor.read_full()
        
        assert isinstance(df, pl.DataFrame)
        assert len(df) == 1000  # 除去标题行
        assert len(df.columns) == 4
        assert 'name' in df.columns
        assert 'age' in df.columns
    
    @pytest.mark.asyncio
    async def test_csv_read_stream(self, sample_csv_file):
        """测试流式CSV读取"""
        processor = CSVProcessor(sample_csv_file, encoding='utf-8', delimiter=',')
        
        chunks = []
        async for chunk in processor.read_stream(chunk_size=200):
            chunks.append(chunk)
            assert isinstance(chunk, pl.DataFrame)
        
        # 应该有多个块
        assert len(chunks) > 1
        
        # 合并后的总行数应该等于原始数据
        total_rows = sum(len(chunk) for chunk in chunks)
        assert total_rows == 1000
    
    @pytest.mark.asyncio
    async def test_csv_get_sample(self, sample_csv_file):
        """测试CSV样本获取"""
        processor = CSVProcessor(sample_csv_file, encoding='utf-8', delimiter=',')
        sample = await processor.get_sample(sample_rows=100)
        
        assert isinstance(sample, pl.DataFrame)
        assert len(sample) <= 100
        assert len(sample.columns) == 4
    
    def test_csv_memory_estimation(self, sample_csv_file):
        """测试内存使用估算"""
        processor = CSVProcessor(sample_csv_file, encoding='utf-8', delimiter=',')
        estimated_memory = processor._estimate_memory_usage()
        
        assert isinstance(estimated_memory, int)
        assert estimated_memory > 0


class TestExcelProcessor:
    """Excel处理器测试"""
    
    @pytest.fixture
    def sample_excel_file(self, tmp_path):
        """创建样本Excel文件"""
        excel_file = tmp_path / "sample.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        
        # 添加标题行
        ws.append(['name', 'age', 'city', 'salary'])
        
        # 添加数据行
        for i in range(1000):
            ws.append([f'Name{i}', 20 + i % 50, f'City{i % 10}', 30000 + i * 100])
        
        wb.save(str(excel_file))
        return str(excel_file)
    
    @pytest.mark.asyncio
    async def test_excel_read_full(self, sample_excel_file):
        """测试完整Excel读取"""
        processor = ExcelProcessor(sample_excel_file)
        df = await processor.read_full()
        
        assert isinstance(df, pl.DataFrame)
        assert len(df) == 1000
        assert len(df.columns) == 4
        assert 'name' in df.columns
    
    @pytest.mark.asyncio
    async def test_excel_read_stream(self, sample_excel_file):
        """测试流式Excel读取"""
        processor = ExcelProcessor(sample_excel_file)
        
        chunks = []
        async for chunk in processor.read_stream(chunk_size=200):
            chunks.append(chunk)
            assert isinstance(chunk, pl.DataFrame)
        
        # 应该有多个块
        assert len(chunks) > 1
        
        # 合并后的总行数应该等于原始数据
        total_rows = sum(len(chunk) for chunk in chunks)
        assert total_rows == 1000
    
    @pytest.mark.asyncio
    async def test_excel_get_sample(self, sample_excel_file):
        """测试Excel样本获取"""
        processor = ExcelProcessor(sample_excel_file)
        sample = await processor.get_sample(sample_rows=100)
        
        assert isinstance(sample, pl.DataFrame)
        assert len(sample) <= 100
    
    @pytest.mark.asyncio
    async def test_excel_sheet_info(self, sample_excel_file):
        """测试工作表信息获取"""
        processor = ExcelProcessor(sample_excel_file)
        sheet_info = await processor.get_sheet_info()
        
        assert isinstance(sheet_info, dict)
        assert 'sheet_names' in sheet_info
        assert 'sheet_count' in sheet_info


class TestFileConnector:
    """文件连接器测试"""
    
    @pytest.fixture
    def sample_csv_file(self, tmp_path):
        """创建样本CSV文件"""
        csv_file = tmp_path / "test.csv"
        data = []
        data.append("id,name,age,city,salary")
        for i in range(10000):  # 1万行测试数据
            data.append(f"{i},Name{i},{20 + i % 50},City{i % 10},{30000 + i * 10}")
        csv_file.write_text("\n".join(data))
        return str(csv_file)
    
    @pytest.fixture
    def sample_excel_file(self, tmp_path):
        """创建样本Excel文件"""
        excel_file = tmp_path / "test.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        
        ws.append(['id', 'name', 'age', 'city', 'salary'])
        for i in range(1000):
            ws.append([i, f'Name{i}', 20 + i % 50, f'City{i % 10}', 30000 + i * 10])
        
        wb.save(str(excel_file))
        return str(excel_file)
    
    @pytest.mark.asyncio
    async def test_csv_connector_connection(self, sample_csv_file):
        """测试CSV连接器连接"""
        connector = FileConnector(
            source_id="test_csv",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        # 测试连接
        connected = await connector.connect()
        assert connected is True
        assert connector.connection_status == ConnectionStatus.CONNECTED
        
        # 测试连接验证
        is_valid = await connector.test_connection()
        assert is_valid is True
        
        # 关闭连接
        await connector.close()
        assert connector.connection_status == ConnectionStatus.CLOSED
    
    @pytest.mark.asyncio
    async def test_excel_connector_connection(self, sample_excel_file):
        """测试Excel连接器连接"""
        connector = FileConnector(
            source_id="test_excel",
            source_type=DataSourceType.EXCEL,
            file_path=sample_excel_file
        )
        
        connected = await connector.connect()
        assert connected is True
        assert connector.connection_status == ConnectionStatus.CONNECTED
        
        is_valid = await connector.test_connection()
        assert is_valid is True
        
        await connector.close()
        assert connector.connection_status == ConnectionStatus.CLOSED
    
    @pytest.mark.asyncio
    async def test_get_schema(self, sample_csv_file):
        """测试schema获取"""
        connector = FileConnector(
            source_id="test_csv",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        await connector.connect()
        schema = await connector.get_schema()
        
        assert isinstance(schema, dict)
        assert 'columns' in schema
        assert 'row_count' in schema
        assert 'column_count' in schema
        assert 'file_info' in schema
        
        assert len(schema['columns']) == 5  # id, name, age, city, salary
        assert schema['column_count'] == 5
        assert schema['row_count'] > 0
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_query_basic(self, sample_csv_file):
        """测试基本查询"""
        connector = FileConnector(
            source_id="test_csv",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        await connector.connect()
        
        # 无过滤查询
        df = await connector.query()
        assert isinstance(df, pl.DataFrame)
        assert len(df) == 10000
        assert len(df.columns) == 5
        
        # 带限制的查询
        options = QueryOptions(limit=100)
        df_limited = await connector.query(options=options)
        assert len(df_limited) == 100
        
        # 带跳过的查询
        options = QueryOptions(skip=100, limit=50)
        df_skip = await connector.query(options=options)
        assert len(df_skip) == 50
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_query_with_filters(self, sample_csv_file):
        """测试带过滤器的查询"""
        connector = FileConnector(
            source_id="test_csv",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        await connector.connect()
        
        # 等值过滤
        query = {"city": "City0"}
        df = await connector.query(query=query)
        assert isinstance(df, pl.DataFrame)
        assert len(df) > 0
        # 验证所有结果都匹配过滤条件
        assert all(city == "City0" for city in df['city'].to_list())
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_query_with_projection(self, sample_csv_file):
        """测试带投影的查询"""
        connector = FileConnector(
            source_id="test_csv",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        await connector.connect()
        
        # 选择特定列
        options = QueryOptions(
            projection={"name": 1, "age": 1},
            limit=100
        )
        df = await connector.query(options=options)
        assert len(df.columns) == 2
        assert 'name' in df.columns
        assert 'age' in df.columns
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_query_stream(self, sample_csv_file):
        """测试流式查询"""
        connector = FileConnector(
            source_id="test_csv",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        await connector.connect()
        
        chunks = []
        options = QueryOptions(batch_size=1000)
        
        async for chunk in connector.query_stream(options=options):
            chunks.append(chunk)
            assert isinstance(chunk, pl.DataFrame)
        
        # 应该有多个块
        assert len(chunks) > 1
        
        # 合并后的总行数应该等于原始数据
        total_rows = sum(len(chunk) for chunk in chunks)
        assert total_rows == 10000
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_get_row_count(self, sample_csv_file):
        """测试行数获取"""
        connector = FileConnector(
            source_id="test_csv",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        await connector.connect()
        
        # 总行数
        total_count = await connector.get_row_count()
        assert total_count == 10000
        
        # 带过滤的行数
        query = {"city": "City0"}
        filtered_count = await connector.get_row_count(query)
        assert filtered_count > 0
        assert filtered_count < total_count
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_preview_data(self, sample_csv_file):
        """测试数据预览"""
        connector = FileConnector(
            source_id="test_csv",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        await connector.connect()
        
        preview = await connector.preview_data(limit=50)
        assert isinstance(preview, pl.DataFrame)
        assert len(preview) <= 50
        assert len(preview.columns) == 5
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_auto_encoding_detection(self, tmp_path):
        """测试自动编码检测"""
        # 创建GBK编码的文件
        gbk_file = tmp_path / "gbk_test.csv"
        content = "姓名,年龄,城市\n张三,25,北京\n李四,30,上海"
        
        with open(str(gbk_file), 'w', encoding='gbk') as f:
            f.write(content)
        
        connector = FileConnector(
            source_id="test_gbk",
            source_type=DataSourceType.CSV,
            file_path=str(gbk_file),
            encoding='auto'  # 自动检测编码
        )
        
        connected = await connector.connect()
        assert connected is True
        
        df = await connector.query()
        assert isinstance(df, pl.DataFrame)
        assert len(df) == 2
        assert '姓名' in df.columns
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_auto_delimiter_detection(self, tmp_path):
        """测试自动分隔符检测"""
        # 创建制表符分隔的文件
        tsv_file = tmp_path / "test.tsv"
        data = []
        data.append("name\tage\tcity")
        data.append("John\t25\tNYC")
        data.append("Jane\t30\tLA")
        tsv_file.write_text("\n".join(data))
        
        connector = FileConnector(
            source_id="test_tsv",
            source_type=DataSourceType.CSV,
            file_path=str(tsv_file),
            delimiter='auto'  # 自动检测分隔符
        )
        
        connected = await connector.connect()
        assert connected is True
        
        df = await connector.query()
        assert isinstance(df, pl.DataFrame)
        assert len(df) == 2
        assert len(df.columns) == 3
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_error_handling(self):
        """测试错误处理"""
        # 测试不存在的文件
        connector = FileConnector(
            source_id="test_error",
            source_type=DataSourceType.CSV,
            file_path="/non/existent/file.csv"
        )
        
        with pytest.raises(Exception):
            await connector.connect()
        
        # 测试类型不匹配
        # 这里需要一个真实的Excel文件来测试
        # 但声明为CSV类型
    
    @pytest.mark.asyncio
    async def test_memory_usage_monitoring(self, sample_csv_file):
        """测试内存使用监控"""
        connector = FileConnector(
            source_id="test_memory",
            source_type=DataSourceType.CSV,
            file_path=sample_csv_file
        )
        
        await connector.connect()
        
        # 记录初始内存状态
        initial_memory = memory_guard.get_memory_stats()
        
        # 执行查询
        df = await connector.query()
        assert len(df) == 10000
        
        # 检查内存状态（应该有合理的增长）
        final_memory = memory_guard.get_memory_stats()
        memory_delta = final_memory.rss_bytes - initial_memory.rss_bytes
        
        # 内存增长应该在合理范围内（不超过500MB）
        assert memory_delta < 500 * 1024 * 1024
        
        await connector.close()


class TestPerformance:
    """性能测试"""
    
    def create_large_csv(self, file_path: str, rows: int = 1000000):
        """创建大型CSV文件用于性能测试"""
        with open(file_path, 'w', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(['id', 'name', 'age', 'city', 'salary', 'department', 'email'])
            
            for i in range(rows):
                writer.writerow([
                    i,
                    f'Employee_{i}',
                    20 + (i % 50),
                    f'City_{i % 100}',
                    30000 + (i % 100000),
                    f'Dept_{i % 20}',
                    f'user{i}@company.com'
                ])
    
    @pytest.mark.asyncio
    @pytest.mark.slow  # 标记为慢测试
    async def test_large_csv_performance(self, tmp_path):
        """测试100万行CSV性能（目标<15秒）"""
        # 创建100万行CSV文件
        large_csv = tmp_path / "large_test.csv"
        print(f"\n创建100万行测试文件: {large_csv}")
        
        start_time = time.time()
        self.create_large_csv(str(large_csv), rows=1000000)
        create_time = time.time() - start_time
        print(f"文件创建耗时: {create_time:.2f}秒")
        
        # 测试连接和读取性能
        connector = FileConnector(
            source_id="large_csv_test",
            source_type=DataSourceType.CSV,
            file_path=str(large_csv)
        )
        
        # 连接测试
        connect_start = time.time()
        await connector.connect()
        connect_time = time.time() - connect_start
        print(f"连接耗时: {connect_time:.2f}秒")
        
        # 读取测试
        read_start = time.time()
        df = await connector.query()
        read_time = time.time() - read_start
        
        print(f"读取100万行耗时: {read_time:.2f}秒")
        print(f"读取速度: {len(df) / read_time:.0f} 行/秒")
        print(f"内存使用: {memory_guard.get_memory_stats().rss_bytes / 1024 / 1024:.1f}MB")
        
        # 验证性能目标
        assert read_time < config.performance.csv_performance_target_seconds, \
            f"读取时间 {read_time:.2f}秒 超过目标 {config.performance.csv_performance_target_seconds}秒"
        
        assert len(df) == 1000000, f"数据行数不匹配: {len(df)} != 1000000"
        assert len(df.columns) == 7, f"列数不匹配: {len(df.columns)} != 7"
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_streaming_performance(self, tmp_path):
        """测试流式处理性能"""
        # 创建测试文件（10万行）
        csv_file = tmp_path / "streaming_test.csv"
        self.create_large_csv(str(csv_file), rows=100000)
        
        connector = FileConnector(
            source_id="streaming_test",
            source_type=DataSourceType.CSV,
            file_path=str(csv_file)
        )
        
        await connector.connect()
        
        # 测试流式处理
        start_time = time.time()
        total_rows = 0
        chunk_count = 0
        
        options = QueryOptions(batch_size=5000, streaming=True)
        async for chunk in connector.query_stream(options=options):
            total_rows += len(chunk)
            chunk_count += 1
        
        stream_time = time.time() - start_time
        
        print(f"\n流式处理10万行:")
        print(f"总耗时: {stream_time:.2f}秒")
        print(f"块数: {chunk_count}")
        print(f"总行数: {total_rows}")
        print(f"处理速度: {total_rows / stream_time:.0f} 行/秒")
        
        assert total_rows == 100000
        assert chunk_count > 1  # 应该分成多个块
        
        await connector.close()
    
    @pytest.mark.asyncio
    async def test_memory_efficiency(self, tmp_path):
        """测试内存效率"""
        # 创建较大的测试文件（50万行）
        csv_file = tmp_path / "memory_test.csv"
        self.create_large_csv(str(csv_file), rows=500000)
        
        connector = FileConnector(
            source_id="memory_test",
            source_type=DataSourceType.CSV,
            file_path=str(csv_file)
        )
        
        await connector.connect()
        
        # 记录初始内存
        initial_memory = memory_guard.get_memory_stats().rss_bytes
        
        # 执行查询
        df = await connector.query()
        
        # 记录峰值内存
        peak_memory = memory_guard.get_memory_stats().rss_bytes
        memory_delta = peak_memory - initial_memory
        
        print(f"\n内存效率测试（50万行）:")
        print(f"初始内存: {initial_memory / 1024 / 1024:.1f}MB")
        print(f"峰值内存: {peak_memory / 1024 / 1024:.1f}MB")
        print(f"内存增长: {memory_delta / 1024 / 1024:.1f}MB")
        print(f"数据行数: {len(df)}")
        print(f"每行内存: {memory_delta / len(df):.1f} 字节/行")
        
        # 验证内存使用合理性
        # 50万行，7列数据，每行应该不超过1KB内存
        max_expected_memory = len(df) * 1000  # 1KB per row
        assert memory_delta < max_expected_memory, \
            f"内存使用过高: {memory_delta / 1024 / 1024:.1f}MB"
        
        await connector.close()


class TestRegistryIntegration:
    """注册系统集成测试"""
    
    @pytest.mark.asyncio
    async def test_connector_registration(self, tmp_path):
        """测试连接器注册"""
        # 创建测试文件
        csv_file = tmp_path / "registry_test.csv"
        csv_file.write_text("name,age\nJohn,25\nJane,30")
        
        # 通过注册系统创建连接器
        connector = await registry.register_source(
            source_id="registry_csv_test",
            source_type=DataSourceType.CSV,
            file_path=str(csv_file)
        )
        
        assert connector is not None
        assert connector.source_id == "registry_csv_test"
        assert connector.source_type == DataSourceType.CSV
        
        # 连接并测试
        await registry.connect_source("registry_csv_test")
        
        # 验证连接状态
        registered_connector = registry.get_source("registry_csv_test")
        assert registered_connector.connection_status == ConnectionStatus.CONNECTED
        
        # 清理
        await registry.unregister_source("registry_csv_test")
    
    @pytest.mark.asyncio
    async def test_multiple_connectors(self, tmp_path):
        """测试多个连接器同时工作"""
        # 创建多个测试文件
        files = []
        for i in range(3):
            csv_file = tmp_path / f"multi_test_{i}.csv"
            data = []
            data.append("id,name,value")
            for j in range(100):
                data.append(f"{j},Name_{i}_{j},{j * 10}")
            csv_file.write_text("\n".join(data))
            files.append(str(csv_file))
        
        # 注册多个连接器
        connectors = []
        for i, file_path in enumerate(files):
            connector = await registry.register_source(
                source_id=f"multi_csv_{i}",
                source_type=DataSourceType.CSV,
                file_path=file_path
            )
            connectors.append(connector)
            await registry.connect_source(f"multi_csv_{i}")
        
        # 测试并发查询
        tasks = []
        for i in range(3):
            connector = registry.get_source(f"multi_csv_{i}")
            task = connector.query()
            tasks.append(task)
        
        results = await asyncio.gather(*tasks)
        
        # 验证结果
        for i, df in enumerate(results):
            assert isinstance(df, pl.DataFrame)
            assert len(df) == 100
            assert f"Name_{i}_0" in df['name'].to_list()
        
        # 清理
        for i in range(3):
            await registry.unregister_source(f"multi_csv_{i}")
    
    def test_connector_type_registration(self):
        """测试连接器类型注册"""
        # 验证CSV和Excel连接器已注册
        csv_sources = registry.list_sources(source_type=DataSourceType.CSV)
        excel_sources = registry.list_sources(source_type=DataSourceType.EXCEL)
        
        # 检查注册信息
        info = asyncio.run(registry.get_registry_info())
        registered_types = info['registered_types']
        
        assert 'csv' in registered_types
        assert 'excel' in registered_types


if __name__ == "__main__":
    # 运行测试
    pytest.main([
        __file__,
        "-v",
        "--tb=short",
        "-m", "not slow",  # 默认跳过慢测试
        "--asyncio-mode=auto"
    ])
    
    # 运行性能测试（可选）
    print("\n=== 运行性能测试 ===")
    pytest.main([
        __file__ + "::TestPerformance",
        "-v",
        "--tb=short",
        "-m", "slow",
        "--asyncio-mode=auto"
    ])
"""
文件处理工具函数
提供文件类型检测、大小估算、路径处理等通用功能
"""
import os
import mimetypes
import tempfile
import shutil
import hashlib
from pathlib import Path
from typing import Optional, Dict, Any, List, Tuple, Union
import logging
import csv
from io import StringIO, BytesIO


class FileUtils:
    """文件处理工具类"""
    
    # 支持的文件类型
    SUPPORTED_EXTENSIONS = {
        '.csv': 'csv',
        '.tsv': 'csv',
        '.txt': 'csv',
        '.xlsx': 'excel',
        '.xls': 'excel',
        '.xlsm': 'excel'
    }
    
    # CSV分隔符检测模式
    CSV_DELIMITERS = [',', '\t', ';', '|', ':']
    
    def __init__(self):
        self._logger = logging.getLogger(self.__class__.__name__)
        
        # 初始化MIME类型
        mimetypes.init()
    
    def get_file_type(self, file_path: str) -> str:
        """检测文件类型
        
        Args:
            file_path: 文件路径
            
        Returns:
            str: 文件类型 ('csv', 'excel', 'unknown')
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        # 1. 通过扩展名检测
        ext = Path(file_path).suffix.lower()
        if ext in self.SUPPORTED_EXTENSIONS:
            return self.SUPPORTED_EXTENSIONS[ext]
        
        # 2. 通过MIME类型检测
        mime_type, _ = mimetypes.guess_type(file_path)
        if mime_type:
            if mime_type in ['text/csv', 'text/plain']:
                return 'csv'
            elif mime_type in [
                'application/vnd.ms-excel',
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            ]:
                return 'excel'
        
        # 3. 通过文件头部检测
        try:
            with open(file_path, 'rb') as f:
                header = f.read(8)
            
            # Excel文件头部标识
            if header.startswith(b'PK\x03\x04'):  # .xlsx
                return 'excel'
            elif header.startswith(b'\xd0\xcf\x11\xe0'):  # .xls
                return 'excel'
            
            # 尝试作为CSV检测
            if self._is_likely_csv(file_path):
                return 'csv'
                
        except Exception as e:
            self._logger.warning(f"File header detection failed: {e}")
        
        return 'unknown'
    
    def _is_likely_csv(self, file_path: str) -> bool:
        """检测文件是否可能是CSV格式"""
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                # 读取前几行
                lines = [f.readline() for _ in range(3)]
            
            # 检查是否包含常见的CSV分隔符
            for line in lines:
                if any(delim in line for delim in self.CSV_DELIMITERS):
                    return True
            
            return False
        except Exception:
            return False
    
    def detect_csv_delimiter(
        self, 
        file_path: str, 
        sample_lines: int = 5,
        encoding: str = 'utf-8'
    ) -> str:
        """检测CSV分隔符
        
        Args:
            file_path: 文件路径
            sample_lines: 采样行数
            encoding: 文件编码
            
        Returns:
            str: 检测到的分隔符
        """
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                sample = ''.join([f.readline() for _ in range(sample_lines)])
            
            # 使用csv.Sniffer检测
            sniffer = csv.Sniffer()
            try:
                dialect = sniffer.sniff(sample, delimiters=',\t;|:')
                return dialect.delimiter
            except csv.Error:
                # Sniffer失败时，手动检测
                return self._manual_delimiter_detection(sample)
                
        except Exception as e:
            self._logger.warning(f"Delimiter detection failed: {e}")
            return ','  # 默认逗号
    
    def _manual_delimiter_detection(self, sample: str) -> str:
        """手动检测分隔符"""
        delimiter_counts = {}
        
        for delimiter in self.CSV_DELIMITERS:
            count = sample.count(delimiter)
            if count > 0:
                # 计算每行的平均分隔符数量
                lines = sample.strip().split('\n')
                if lines:
                    avg_count = count / len(lines)
                    delimiter_counts[delimiter] = avg_count
        
        if delimiter_counts:
            # 选择平均数量最多的分隔符
            return max(delimiter_counts, key=delimiter_counts.get)
        
        return ','  # 默认返回逗号
    
    def estimate_file_info(self, file_path: str) -> Dict[str, Any]:
        """估算文件信息
        
        Args:
            file_path: 文件路径
            
        Returns:
            Dict[str, Any]: 文件信息
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_stats = os.stat(file_path)
        file_info = {
            'path': file_path,
            'name': os.path.basename(file_path),
            'size_bytes': file_stats.st_size,
            'size_mb': file_stats.st_size / (1024 * 1024),
            'modified_time': file_stats.st_mtime,
            'file_type': self.get_file_type(file_path)
        }
        
        # 估算行数（仅对文本文件）
        if file_info['file_type'] == 'csv':
            file_info.update(self._estimate_csv_info(file_path))
        elif file_info['file_type'] == 'excel':
            file_info.update(self._estimate_excel_info(file_path))
        
        return file_info
    
    def _estimate_csv_info(self, file_path: str) -> Dict[str, Any]:
        """估算CSV文件信息"""
        info = {}
        
        try:
            # 采样前1000行估算
            line_count = 0
            total_length = 0
            encoding = 'utf-8'
            
            # 简单编码检测
            with open(file_path, 'rb') as f:
                sample = f.read(1000)
                try:
                    sample.decode('utf-8')
                except UnicodeDecodeError:
                    encoding = 'latin-1'
            
            with open(file_path, 'r', encoding=encoding, errors='ignore') as f:
                for i, line in enumerate(f):
                    if i >= 1000:  # 限制采样
                        break
                    line_count += 1
                    total_length += len(line)
            
            if line_count > 0:
                avg_line_length = total_length / line_count
                file_size = os.path.getsize(file_path)
                estimated_rows = int(file_size / avg_line_length)
                
                info.update({
                    'estimated_rows': max(1, estimated_rows),
                    'avg_line_length': avg_line_length,
                    'sample_lines': line_count,
                    'delimiter': self.detect_csv_delimiter(file_path, encoding=encoding)
                })
            
        except Exception as e:
            self._logger.warning(f"CSV info estimation failed: {e}")
            info.update({
                'estimated_rows': 0,
                'avg_line_length': 0,
                'sample_lines': 0,
                'delimiter': ','
            })
        
        return info
    
    def _estimate_excel_info(self, file_path: str) -> Dict[str, Any]:
        """估算Excel文件信息"""
        info = {
            'estimated_rows': 0,
            'estimated_columns': 0,
            'sheet_count': 0,
            'sheet_names': []
        }
        
        try:
            import openpyxl
            from openpyxl import load_workbook
            
            # 只读模式加载工作簿
            wb = load_workbook(file_path, read_only=True, data_only=True)
            info['sheet_names'] = wb.sheetnames
            info['sheet_count'] = len(wb.sheetnames)
            
            # 获取第一个工作表的信息
            if wb.worksheets:
                ws = wb.worksheets[0]
                info['estimated_rows'] = ws.max_row
                info['estimated_columns'] = ws.max_column
            
            wb.close()
            
        except ImportError:
            self._logger.warning("openpyxl not available for Excel info estimation")
        except Exception as e:
            self._logger.warning(f"Excel info estimation failed: {e}")
        
        return info
    
    def validate_file_access(self, file_path: str) -> Dict[str, Any]:
        """验证文件访问权限
        
        Args:
            file_path: 文件路径
            
        Returns:
            Dict[str, Any]: 验证结果
        """
        result = {
            'exists': False,
            'readable': False,
            'size_ok': False,
            'type_supported': False,
            'errors': []
        }
        
        try:
            # 检查文件是否存在
            if not os.path.exists(file_path):
                result['errors'].append(f"File does not exist: {file_path}")
                return result
            result['exists'] = True
            
            # 检查是否可读
            if not os.access(file_path, os.R_OK):
                result['errors'].append(f"File is not readable: {file_path}")
                return result
            result['readable'] = True
            
            # 检查文件大小
            file_size = os.path.getsize(file_path)
            max_size = 2 * 1024 * 1024 * 1024  # 2GB限制
            if file_size > max_size:
                result['errors'].append(f"File too large: {file_size / 1024 / 1024:.1f}MB > 2GB")
                return result
            result['size_ok'] = True
            
            # 检查文件类型
            file_type = self.get_file_type(file_path)
            if file_type == 'unknown':
                result['errors'].append(f"Unsupported file type: {file_path}")
                return result
            result['type_supported'] = True
            result['file_type'] = file_type
            
        except Exception as e:
            result['errors'].append(f"Validation error: {str(e)}")
        
        return result
    
    def create_temp_file(
        self, 
        content: Union[str, bytes], 
        suffix: str = '.tmp',
        prefix: str = 'file_connector_'
    ) -> str:
        """创建临时文件
        
        Args:
            content: 文件内容
            suffix: 文件后缀
            prefix: 文件前缀
            
        Returns:
            str: 临时文件路径
        """
        fd, temp_path = tempfile.mkstemp(suffix=suffix, prefix=prefix)
        
        try:
            if isinstance(content, str):
                with os.fdopen(fd, 'w', encoding='utf-8') as f:
                    f.write(content)
            else:
                with os.fdopen(fd, 'wb') as f:
                    f.write(content)
        except Exception:
            # 如果写入失败，清理文件描述符和文件
            try:
                os.close(fd)
            except:
                pass
            if os.path.exists(temp_path):
                os.unlink(temp_path)
            raise
        
        return temp_path
    
    def cleanup_temp_file(self, temp_path: str) -> bool:
        """清理临时文件
        
        Args:
            temp_path: 临时文件路径
            
        Returns:
            bool: 清理是否成功
        """
        try:
            if os.path.exists(temp_path):
                os.unlink(temp_path)
            return True
        except Exception as e:
            self._logger.error(f"Failed to cleanup temp file {temp_path}: {e}")
            return False
    
    def get_file_hash(self, file_path: str, algorithm: str = 'md5') -> str:
        """计算文件哈希值
        
        Args:
            file_path: 文件路径
            algorithm: 哈希算法 ('md5', 'sha256')
            
        Returns:
            str: 文件哈希值
        """
        if algorithm == 'md5':
            hash_func = hashlib.md5()
        elif algorithm == 'sha256':
            hash_func = hashlib.sha256()
        else:
            raise ValueError(f"Unsupported hash algorithm: {algorithm}")
        
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(8192), b""):
                hash_func.update(chunk)
        
        return hash_func.hexdigest()
    
    def normalize_path(self, path: str) -> str:
        """标准化文件路径
        
        Args:
            path: 原始路径
            
        Returns:
            str: 标准化后的绝对路径
        """
        # 展开用户目录
        path = os.path.expanduser(path)
        # 展开环境变量
        path = os.path.expandvars(path)
        # 转换为绝对路径
        path = os.path.abspath(path)
        # 标准化路径
        path = os.path.normpath(path)
        
        return path
    
    def safe_filename(self, filename: str) -> str:
        """生成安全的文件名
        
        Args:
            filename: 原始文件名
            
        Returns:
            str: 安全的文件名
        """
        # 移除或替换危险字符
        unsafe_chars = '<>:"/\\|?*'
        safe_name = filename
        
        for char in unsafe_chars:
            safe_name = safe_name.replace(char, '_')
        
        # 限制长度
        if len(safe_name) > 255:
            name, ext = os.path.splitext(safe_name)
            max_name_len = 255 - len(ext)
            safe_name = name[:max_name_len] + ext
        
        # 避免保留名称
        reserved_names = [
            'CON', 'PRN', 'AUX', 'NUL',
            'COM1', 'COM2', 'COM3', 'COM4', 'COM5', 'COM6', 'COM7', 'COM8', 'COM9',
            'LPT1', 'LPT2', 'LPT3', 'LPT4', 'LPT5', 'LPT6', 'LPT7', 'LPT8', 'LPT9'
        ]
        
        name_only = os.path.splitext(safe_name)[0].upper()
        if name_only in reserved_names:
            safe_name = f"_{safe_name}"
        
        return safe_name
    
    def copy_file_with_progress(
        self, 
        src: str, 
        dst: str, 
        callback: Optional[callable] = None
    ) -> bool:
        """带进度的文件复制
        
        Args:
            src: 源文件路径
            dst: 目标文件路径
            callback: 进度回调函数，参数为 (copied_bytes, total_bytes)
            
        Returns:
            bool: 复制是否成功
        """
        try:
            if not os.path.exists(src):
                raise FileNotFoundError(f"Source file not found: {src}")
            
            # 确保目标目录存在
            dst_dir = os.path.dirname(dst)
            if dst_dir and not os.path.exists(dst_dir):
                os.makedirs(dst_dir, exist_ok=True)
            
            file_size = os.path.getsize(src)
            copied_bytes = 0
            
            with open(src, 'rb') as src_file, open(dst, 'wb') as dst_file:
                while True:
                    chunk = src_file.read(8192)  # 8KB chunks
                    if not chunk:
                        break
                    
                    dst_file.write(chunk)
                    copied_bytes += len(chunk)
                    
                    if callback:
                        callback(copied_bytes, file_size)
            
            return True
            
        except Exception as e:
            self._logger.error(f"File copy failed: {e}")
            return False


# 全局文件工具实例
file_utils = FileUtils()